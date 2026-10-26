"""Text extraction for uploaded documents."""

from __future__ import annotations

import io

ALLOWED_TYPES: dict[str, str] = {
    "application/pdf": "pdf",
    "application/vnd.openxmlformats-officedocument.wordprocessingml.document": "docx",
    "text/plain": "txt",
    "text/markdown": "md",
}
_EXT_FALLBACK = {".pdf": "pdf", ".docx": "docx", ".txt": "txt", ".md": "md"}


class UnsupportedDocument(Exception):
    pass


def detect_kind(filename: str, content_type: str) -> str:
    kind = ALLOWED_TYPES.get(content_type)
    if kind:
        return kind
    for ext, k in _EXT_FALLBACK.items():
        if filename.lower().endswith(ext):
            return k
    raise UnsupportedDocument(f"Unsupported file type: {content_type or filename}")


def extract_text(data: bytes, kind: str) -> str:
    if kind == "pdf":
        from pypdf import PdfReader

        reader = PdfReader(io.BytesIO(data))
        if reader.is_encrypted:
            raise UnsupportedDocument("Encrypted PDFs are not supported")
        return "\n\n".join((page.extract_text() or "") for page in reader.pages)
    if kind == "docx":
        from docx import Document as DocxDocument

        doc = DocxDocument(io.BytesIO(data))
        parts = [p.text for p in doc.paragraphs if p.text.strip()]
        for table in doc.tables:
            for row in table.rows:
                parts.append(" | ".join(cell.text.strip() for cell in row.cells))
        return "\n\n".join(parts)
    if kind in ("txt", "md"):
        return data.decode("utf-8", errors="replace")
    raise UnsupportedDocument(f"Unsupported kind: {kind}")
